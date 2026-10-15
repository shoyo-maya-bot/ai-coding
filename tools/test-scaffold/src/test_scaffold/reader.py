"""Excel（.xlsx）のテスト仕様書を読み、テストケースの行にする.

- 結合セルは、結合範囲のすべてのセルに左上の値があるものとして読む（縦に結合した「テスト対象クラス」など）
- 空行（対応づけた列にそのセル自身の値が 1 つも無い行）は飛ばす。結合で埋めた値だけの行も空行とみなす
- 複数シートを順に読む（設定の sheets で絞れる）。見出し行の無いシート（表紙・改訂履歴など）は飛ばす
"""

from __future__ import annotations

import datetime as dt
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from .config import REQUIRED_COLUMNS, Config, column_index, normalize_header
from .errors import UsageError

_REQ_SPLIT = re.compile(r"[\s,、，;；/／]+")


@dataclass
class Case:
    sheet: str
    row: int
    location: str  # 例 会員登録!A5（項番のセル）
    id: str
    target_class: str
    target_method: str = ""
    requirement_ids: list[str] = field(default_factory=list)
    viewpoint: str = ""
    case_type: str = ""  # 対応表の proofs の type（normal / abnormal / boundary）。分からなければ空
    precondition: str = ""
    input: str = ""
    expected: str = ""


@dataclass
class Notice:
    location: str
    message: str


@dataclass
class ReadResult:
    cases: list[Case] = field(default_factory=list)
    warnings: list[Notice] = field(default_factory=list)
    skipped_sheets: list[Notice] = field(default_factory=list)


def cell_text(value: object) -> str:
    """セルの値を文字列に（1.0 → '1'、日付は ISO 形式、改行は LF）."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


class _Sheet:
    """結合セルを展開して読むための薄い包み."""

    def __init__(self, ws: Worksheet) -> None:
        self.ws = ws
        self.fill: dict[tuple[int, int], object] = {}
        for rng in ws.merged_cells.ranges:
            top_left = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        self.fill[(r, c)] = top_left

    def own(self, row: int, col: int) -> str:
        """セル自身の値（結合で埋めた値は含めない）."""
        if (row, col) in self.fill:
            return ""
        return cell_text(self.ws.cell(row, col).value)

    def value(self, row: int, col: int) -> str:
        if (row, col) in self.fill:
            return cell_text(self.fill[(row, col)])
        return cell_text(self.ws.cell(row, col).value)


def _map_columns(sheet: _Sheet, row: int, cfg: Config) -> dict[str, int]:
    """見出し行から 論理名 → 列番号 を作る（列記号で指定した列はそのまま）."""
    found: dict[str, int] = {}
    headers = {}
    for col in range(1, (sheet.ws.max_column or 0) + 1):
        key = normalize_header(sheet.value(row, col))
        if key and key not in headers:
            headers[key] = col
    for name, spec in cfg.columns.items():
        if spec.letter:
            found[name] = column_index(spec.letter)
            continue
        for h in spec.headers:
            col = headers.get(normalize_header(h))
            if col is not None:
                found[name] = col
                break
    return found


def _has_required(mapping: dict[str, int]) -> bool:
    return all(name in mapping for name in REQUIRED_COLUMNS)


def _find_header(sheet: _Sheet, cfg: Config) -> tuple[int, dict[str, int]] | None:
    rows = [cfg.header_row] if cfg.header_row else range(1, min(cfg.header_search_rows, sheet.ws.max_row or 0) + 1)
    for r in rows:
        mapping = _map_columns(sheet, r, cfg)
        if _has_required(mapping):
            return r, mapping
    return None


def _missing_columns(sheet: _Sheet, cfg: Config) -> list[str]:
    row = cfg.header_row or 1
    mapping = _map_columns(sheet, row, cfg)
    return [n for n in REQUIRED_COLUMNS if n not in mapping]


def _case_type(raw: str, viewpoint: str, cfg: Config) -> str:
    key = normalize_header(raw)
    if key:
        return cfg.case_types.get(key, "")
    # 種別の列が無いときは、観点の先頭（「異常系: …」など）で判断する
    vp = normalize_header(viewpoint)
    for k, v in cfg.case_types.items():
        if k and vp.startswith(k):
            return v
    return ""


def read_workbook(path: Path, cfg: Config) -> ReadResult:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise UsageError(f"入力は .xlsx / .xlsm だけです（.xls は Excel で .xlsx に保存し直してください）: {path.name}")
    if not path.is_file():
        raise UsageError(f"テスト仕様書がありません: {path}")
    try:
        wb = load_workbook(path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as e:
        raise UsageError(f"Excel として読めません: {path.name}: {e}") from e

    names = [ws.title for ws in wb.worksheets]
    if cfg.sheets:
        missing = [s for s in cfg.sheets if s not in names]
        if missing:
            raise UsageError(f"シートがありません: {', '.join(missing)}（あるシート: {', '.join(names)}）")
        targets = [wb[s] for s in cfg.sheets]
    else:
        targets = list(wb.worksheets)

    result = ReadResult()
    for ws in targets:
        sheet = _Sheet(ws)
        header = _find_header(sheet, cfg)
        if header is None:
            if cfg.sheets:  # 明示したシートに見出しが無いのは設定の誤り
                where = f"{cfg.header_row} 行目" if cfg.header_row else f"先頭 {cfg.header_search_rows} 行"
                raise UsageError(
                    f"シート {ws.title} の{where}に見出しがありません（必要な列: "
                    f"{', '.join(_missing_columns(sheet, cfg)) or ', '.join(REQUIRED_COLUMNS)}）"
                )
            result.skipped_sheets.append(Notice(ws.title, "見出し行が見つからないため飛ばしました"))
            continue
        header_row, mapping = header
        _read_rows(sheet, header_row, mapping, cfg, result)

    if len(result.skipped_sheets) == len(targets):
        raise UsageError(
            "どのシートにも見出し行が見つかりません。header_row と columns を確かめてください"
            f"（必要な列: {', '.join(REQUIRED_COLUMNS)}。--print-config で既定値を表示）"
        )
    return result


def _read_rows(sheet: _Sheet, header_row: int, mapping: dict[str, int], cfg: Config, result: ReadResult) -> None:
    ws = sheet.ws
    for r in range(header_row + 1, (ws.max_row or 0) + 1):
        if not any(sheet.own(r, c) for c in mapping.values()):
            continue  # 空行（結合で埋まった値だけの行も含む）
        get = {name: sheet.value(r, col) for name, col in mapping.items()}
        loc = f"{ws.title}!{get_column_letter(mapping['id'])}{r}"
        case_id = get.get("id", "").replace("\n", " ")
        if not case_id:
            result.warnings.append(Notice(loc, "項番が空のため、この行は飛ばしました"))
            continue
        target = get.get("target_class", "")
        method = get.get("target_method", "")
        if "#" in target and not method:
            target, _, method = target.partition("#")
        target, method = target.strip(), method.strip()
        if not target:
            result.warnings.append(Notice(loc, f"項番 {case_id}: テスト対象クラスが空のため、この行は飛ばしました"))
            continue
        req_ids = [x for x in _REQ_SPLIT.split(get.get("requirement_id", "")) if x]
        if not req_ids:
            result.warnings.append(
                Notice(loc, f"項番 {case_id}: 要件 ID がありません（対応表の proofs に結び付けられない）")
            )
        for rid in req_ids:
            if not cfg.requirement_id_pattern.match(rid):
                result.warnings.append(
                    Notice(
                        loc, f"項番 {case_id}: 要件 ID の形式が違います: {rid}（{cfg.requirement_id_pattern.pattern}）"
                    )
                )
        viewpoint = get.get("viewpoint", "")
        if not viewpoint:
            result.warnings.append(Notice(loc, f"項番 {case_id}: 観点が空です"))
        result.cases.append(
            Case(
                sheet=ws.title,
                row=r,
                location=loc,
                id=case_id,
                target_class=target,
                target_method=method,
                requirement_ids=req_ids,
                viewpoint=viewpoint,
                case_type=_case_type(get.get("case_type", ""), viewpoint, cfg),
                precondition=get.get("precondition", ""),
                input=get.get("input", ""),
                expected=get.get("expected", ""),
            )
        )

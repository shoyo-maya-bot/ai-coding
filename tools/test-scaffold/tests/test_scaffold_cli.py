"""test-scaffold のテスト. テスト仕様書（架空の会員・注文）はテストの中で作る."""

from __future__ import annotations

import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from test_scaffold.cli import run_main
from test_scaffold.java import java_len

HEADERS = ["項番", "テスト対象クラス", "テスト対象メソッド", "要件 ID", "種別", "観点", "前提条件", "入力", "期待結果"]
# Checkstyle（template/config/checkstyle/checkstyle.xml）の MethodName・TypeName
METHOD_NAME = re.compile(r"^[a-z][a-zA-Z0-9]*(_[a-zA-Z0-9]+)*$")
MEMBER_TEST = "template/src/test/java/com/example/app/member/MemberServiceTest.java"
ORDER_TEST = "template/src/test/java/com/example/app/order/OrderImportServiceTest.java"


def member_rows() -> list[list[object]]:
    return [
        [1, "MemberService", "register", "REQ-API0002-003", "正常系", "前後の空白を除き、メールを小文字にして登録する",
         "会員が 0 件", "名前「 テスト 太郎 」、メール「 Taro@Example.COM 」", "ID 1 で登録され、メールは小文字"],
        [2, None, "register", "REQ-API0002-002", "異常系", "メールの重複は大文字小文字を区別せずエラー",
         "a@example.com が登録済み", "メール「A@EXAMPLE.com」", "BusinessException（APPE0003）"],
        [3, None, "get", "REQ-API0001-002", "異常系", "存在しない ID は該当なし", "会員が 0 件", "ID 999",
         "NotFoundException（APPE0002）"],
    ]  # fmt: skip


def make_spec(path: Path, *, extra_member_rows: list[list[object]] | None = None, order: bool = True) -> Path:
    """架空のテスト仕様書: 表紙（見出しなし）・会員（3 行目が見出し、テスト対象クラスを縦に結合、空行あり）・注文."""
    wb = Workbook()
    cover = wb.active
    cover.title = "表紙"
    cover["A1"] = "単体テスト仕様書（サンプル・架空）"
    cover["A3"] = "改訂履歴"

    ws = wb.create_sheet("会員")
    ws["A1"] = "会員 API 単体テスト仕様書"
    ws.merge_cells("A1:I1")
    for col, h in enumerate(HEADERS, start=1):
        ws.cell(3, col, h)
    rows = member_rows()
    for r, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    ws.merge_cells("B4:B6")  # テスト対象クラスを 3 行分結合
    # 7 行目は空行。8 行目以降に追加の行
    for r, row in enumerate(extra_member_rows or [], start=8):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)

    if order:
        wo = wb.create_sheet("注文")
        for col, h in enumerate(HEADERS, start=1):
            wo.cell(1, col, h)
        wo.append(
            ["1-1", "com.example.app.order.OrderImportService", "importLines", "REQ-JOB0001-003", "正常系",
             "登録済みの注文番号は登録しない", "注文 A-001 が登録済み", "A-001 と A-002 の 2 行", "A-002 だけ登録"]
        )  # fmt: skip
        wo.append(
            ["1-2", "com.example.app.order.OrderImportService", "importLines", "REQ-JOB0001-003", "境界値",
             "同じファイルを 2 回取り込んでも件数が変わらない", "", "同じファイルを 2 回", "2 回目の登録 0 件"]
        )  # fmt: skip
    wb.save(path)
    return path


@pytest.fixture
def repo(tmp_path: Path) -> Path:
    """リポジトリのルートに見立てたフォルダ（MemberService の実装クラスがある）."""
    main = tmp_path / "template/src/main/java/com/example/app/member"
    main.mkdir(parents=True)
    (main / "MemberService.java").write_text("package com.example.app.member;\nclass MemberService {}\n", "utf-8")
    return tmp_path


def run(repo: Path, *args: str) -> int:
    return run_main([*args, "--root", str(repo)])


def read(repo: Path, rel: str) -> str:
    return (repo / rel).read_text(encoding="utf-8")


# --- 生成の中身 -----------------------------------------------------------------


def test_generates_one_class_per_target_with_display_name_comments_and_fail(repo, capsys):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec)) == 0
    text = read(repo, MEMBER_TEST)
    assert text.startswith("package com.example.app.member;\n")  # パッケージは実装クラスの場所から
    assert "import static org.junit.jupiter.api.Assertions.fail;" in text
    assert "class MemberServiceTest {" in text
    assert '@DisplayName("1 前後の空白を除き、メールを小文字にして登録する")' in text
    assert "void register_case1() {" in text
    assert "void register_case2() {" in text
    assert "void get_case3() {" in text
    assert 'fail("未実装: 1");' in text
    assert "// 要件 ID: REQ-API0002-003" in text
    assert "// 前提: 会員が 0 件" in text
    assert "// 期待結果: BusinessException（APPE0003）" in text
    assert "// 項番: 2（シート: 会員）" in text
    assert "@Disabled" not in text  # 無効化はしない（guidelines の禁止事項）
    assert str(repo) not in text  # 手元の絶対パスを書き込まない

    order = read(repo, ORDER_TEST)
    assert order.startswith("package com.example.app.order;\n")  # 完全修飾名から
    assert "void importLines_case1_1() {" in order
    assert '@DisplayName("1-2 同じファイルを 2 回取り込んでも件数が変わらない")' in order
    out = capsys.readouterr().out
    assert "作成" in out and "MemberServiceTest#register_case1" in out
    assert "表紙" in out  # 見出しの無いシートは飛ばしたと表示


def test_generated_source_follows_checkstyle_shape(repo):
    long_text = "とても長い観点" * 40 + ' "引用" と \\ を含む'
    spec = make_spec(
        repo / "spec.xlsx",
        extra_member_rows=[
            [4, "MemberService", "register", "REQ-API0002-001", "異常系", long_text, "", "a\tb", "x" * 300]
        ],
    )
    assert run(repo, str(spec)) == 0
    text = read(repo, MEMBER_TEST)
    assert "\t" not in text and "\r" not in text and text.endswith("}\n")
    for line in text.splitlines():
        assert java_len(line) <= 120, line
    for name in re.findall(r"void (\w+)\(", text):
        assert METHOD_NAME.match(name), name
    assert '\\"引用\\"' in text  # 文字列はエスケープする


def test_cell_text_cannot_break_out_of_comments_or_strings(repo):
    evil = 'x\\u000a fail("hacked"); /* */ \\" end'
    spec = make_spec(
        repo / "spec.xlsx", order=False, extra_member_rows=[[5, "MemberService", "get", "", "", evil, evil, evil, evil]]
    )
    assert run(repo, str(spec)) == 0
    text = read(repo, MEMBER_TEST)
    # Java はコメントの中の \u000a も改行として解釈する. バックスラッシュを二重にして無効にする
    assert re.search(r"(?<!\\)(\\\\)*\\u000a", text) is None
    for line in text.splitlines():
        if 'fail("hacked")' in line:
            assert line.lstrip().startswith(("//", '@DisplayName("')) or '+ "' in line


def test_spec_trace_can_read_generated_methods(repo):
    from spec_trace import code  # このリポジトリでは spec-trace も入れる（CI・verify-gates）

    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec), "--proofs-out", "build/proofs.yaml") == 0
    proofs = yaml.safe_load(read(repo, "build/proofs.yaml"))
    entry = proofs["REQ-API0002-003"][0]
    assert entry == {"file": MEMBER_TEST, "symbol": "MemberServiceTest#register_case1", "type": "normal"}
    assert proofs["REQ-JOB0001-003"][1]["type"] == "boundary"
    found = code.lookup_symbol(repo / entry["file"], entry["symbol"])
    assert found.found
    # spec-trace は fail( をアサーションに数える。雛形が proven を通らないのは
    # D3（テストの失敗）と D7 の --test-results（trace.proof_failed）による
    assert found.asserts == 1
    # Gradle の JUnit XML は @DisplayName の表示名を結果の名前にする. spec-trace は表示名でも結果を照合する
    shown = code.display_name(repo / entry["file"], entry["symbol"])
    assert shown == "1 前後の空白を除き、メールを小文字にして登録する"


# --- Excel の読み方 -------------------------------------------------------------


def test_merged_cells_empty_rows_and_sheets(repo, capsys):
    spec = make_spec(repo / "spec.xlsx", extra_member_rows=[[None, None, None, None, None, None, None, None, None]])
    assert run(repo, str(spec), "--dry-run") == 0
    out = capsys.readouterr().out
    # B4:B6 を結合した MemberService が 3 行とも対象クラスになる
    assert out.count("MemberServiceTest#") == 3
    assert out.count("OrderImportServiceTest#") == 2
    assert "警告" not in out


def test_vertical_merge_does_not_turn_blank_rows_into_cases(repo, capsys):
    path = make_spec(repo / "spec.xlsx", order=False)
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb["会員"].unmerge_cells("B4:B6")
    wb["会員"].merge_cells("B4:B9")  # 結合が空行（7〜9 行目）まで伸びている
    wb.save(path)
    assert run(repo, str(path), "--dry-run") == 0
    assert capsys.readouterr().out.count("MemberServiceTest#") == 3


def test_header_row_and_column_letters_from_config(repo, capsys):
    wb = Workbook()
    ws = wb.active
    ws.title = "注文"
    ws["A5"], ws["C5"], ws["D5"], ws["E5"] = "No", "クラス", "確認内容", "Req"
    ws.append(["7", "", "com.example.app.order.OrderImportJob", "引数なしは終了コード 8", "REQ-JOB0001-004"])
    spec = repo / "custom.xlsx"
    wb.save(spec)
    cfg = repo / "scaffold.yaml"
    cfg.write_text(
        "header_row: 5\ncolumns:\n  id: ['No']\n  target_class: {column: C}\n  viewpoint: 確認内容\n"
        "  requirement_id: Req\n",
        encoding="utf-8",
    )
    assert run(repo, str(spec), "-c", str(cfg)) == 0
    text = read(repo, "template/src/test/java/com/example/app/order/OrderImportJobTest.java")
    assert "void case7() {" in text  # 対象メソッドの列が無ければ case<項番>
    assert "// 要件 ID: REQ-JOB0001-004" in text


def test_java_identifiers_are_safe_and_japanese_goes_to_display_name(repo, capsys):
    spec = make_spec(
        repo / "spec.xlsx",
        order=False,
        extra_member_rows=[
            ["No.10-a", "MemberService", "会員を登録する", "REQ-API0002-001", "", "日本語のメソッド名", "", "", ""],
            ["十一", "MemberService", "register(String, String)", "REQ-API0002-001", "", "項番が日本語", "", "", ""],
            [12, "会員サービス", "register", "REQ-API0002-001", "", "日本語のクラス名", "", "", ""],
            [13, "com.example.app.member.MemberService", "1st", "REQ-API0002-001", "", "数字で始まる", "", "", ""],
        ],
    )
    assert run(repo, str(spec)) == 0
    text = read(repo, MEMBER_TEST)
    assert "void caseNo_10_a() {" in text
    assert '@DisplayName("No.10-a 日本語のメソッド名")' in text
    assert "void register_caseR9() {" in text  # 英数字の無い項番は行番号から
    assert "void case13() {" in text
    out = capsys.readouterr().out
    assert "会員サービス" in out and "飛ばしました" in out  # クラス名にできないものは警告して飛ばす
    for name in re.findall(r"void (\w+)\(", text):
        assert METHOD_NAME.match(name), name


# --- 既存ファイル・差分モード・上書き ------------------------------------------------


def test_existing_file_is_not_overwritten_by_default(repo, capsys):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec)) == 0
    target = repo / MEMBER_TEST
    target.write_text("package com.example.app.member;\n\nclass MemberServiceTest {\n    // 人が書いた\n}\n", "utf-8")
    capsys.readouterr()
    assert run(repo, str(spec)) == 0
    assert "人が書いた" in target.read_text("utf-8")
    assert "既にあるため作りませんでした" in capsys.readouterr().out
    assert run(repo, str(spec), "--strict") == 1  # --strict では警告として 1


def test_merge_adds_only_missing_methods_and_keeps_written_tests(repo, capsys):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec)) == 0
    target = repo / MEMBER_TEST
    # 人が項番 1 を実装し（目印のコメントも消した）、項番 2 はメソッド名を変えて目印だけ残した
    text = target.read_text("utf-8")
    text = text.replace('fail("未実装: 1");', "assertThat(1).isEqualTo(1); // 実装済み")
    text = text.replace("// 項番: 1（シート: 会員）\n", "")
    text = text.replace("void register_case2()", "void register_rejectsDuplicateEmail()")
    target.write_text(text, "utf-8")

    spec2 = make_spec(
        repo / "spec2.xlsx",
        extra_member_rows=[
            [4, "MemberService", "register", "REQ-API0002-001", "異常系", "名前が空はエラー", "", "", ""]
        ],
    )
    capsys.readouterr()
    assert run(repo, str(spec2), "--merge") == 0
    merged = target.read_text("utf-8")
    assert "// 実装済み" in merged
    assert merged.count("void register_case1()") == 1
    assert "void register_case2()" not in merged  # 目印（項番 2・シート 会員）があるので足さない
    assert merged.count("void register_case4()") == 1
    assert merged.rstrip().endswith("}")
    assert merged.index("void register_case4()") > merged.index("void get_case3()")
    assert "追加" in capsys.readouterr().out

    before = target.read_text("utf-8")
    assert run(repo, str(spec2), "--merge") == 0  # 2 回目は何も足さない
    assert target.read_text("utf-8") == before
    assert "変更なし" in capsys.readouterr().out


def test_merge_adds_missing_imports_without_clashing_with_assertj_fail(repo):
    spec = make_spec(repo / "spec.xlsx", order=False)
    target = repo / MEMBER_TEST
    target.parent.mkdir(parents=True)
    target.write_text(
        "package com.example.app.member;\n\nimport static org.assertj.core.api.Assertions.fail;\n\n"
        "class MemberServiceTest {\n}\n",
        "utf-8",
    )
    assert run(repo, str(spec), "--merge") == 0
    text = target.read_text("utf-8")
    assert "org.junit.jupiter.api.Assertions.fail" not in text  # 同名の static import を重ねない
    assert "import org.junit.jupiter.api.Test;" in text
    assert "import org.junit.jupiter.api.DisplayName;" in text


def test_force_overwrites(repo):
    spec = make_spec(repo / "spec.xlsx")
    target = repo / MEMBER_TEST
    target.parent.mkdir(parents=True)
    target.write_text("class Old {}\n", "utf-8")
    assert run(repo, str(spec), "--force") == 0
    assert "class MemberServiceTest {" in target.read_text("utf-8")


def test_merge_and_force_are_exclusive(repo):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec), "--merge", "--force") == 2


def test_dry_run_writes_nothing(repo, capsys):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec), "--dry-run", "--proofs-out", "build/proofs.yaml") == 0
    assert not (repo / "template/src/test").exists()
    assert not (repo / "build").exists()
    out = capsys.readouterr().out
    assert "作成予定" in out and MEMBER_TEST in out


# --- パスの検証 -----------------------------------------------------------------


@pytest.mark.parametrize("out", ["../outside", "template/../../outside"])
def test_out_outside_root_is_rejected(repo, out, capsys):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec), "--out", out) == 2
    assert "外を指しています" in capsys.readouterr().err
    assert not (repo.parent / "outside").exists()


def test_absolute_out_outside_root_is_rejected(repo, tmp_path_factory):
    other = tmp_path_factory.mktemp("other")
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec), "--out", str(other)) == 2
    assert list(other.iterdir()) == []


def test_proofs_out_outside_root_is_rejected(repo):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec), "--proofs-out", "../proofs.yaml") == 2
    assert not (repo.parent / "proofs.yaml").exists()


@pytest.mark.skipif(sys.platform == "win32", reason="シンボリックリンクの作成に権限が要る")
def test_symlink_inside_out_pointing_outside_is_rejected(repo, tmp_path_factory):
    other = tmp_path_factory.mktemp("other")
    base = repo / "template/src/test/java/com/example/app"
    base.mkdir(parents=True)
    os.symlink(other, base / "member", target_is_directory=True)
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec)) == 2
    assert list(other.iterdir()) == []


# --- 警告と終了コード ---------------------------------------------------------------


def test_missing_requirement_id_is_warning_and_strict_makes_it_1(repo, capsys):
    spec = make_spec(
        repo / "spec.xlsx",
        order=False,
        extra_member_rows=[[4, "MemberService", "register", "", "異常系", "要件 ID の無い行", "", "", ""]],
    )
    assert run(repo, str(spec), "--dry-run") == 0
    assert "要件 ID がありません" in capsys.readouterr().out
    assert run(repo, str(spec), "--dry-run", "--strict") == 1


def test_malformed_requirement_id_and_unknown_in_trace(repo, capsys):
    trace = repo / "traceability/specs"
    trace.mkdir(parents=True)
    (trace / "member.yaml").write_text(
        "spec_id: FEAT-API0001\nsource: x.md\nrequirements:\n"
        "  - {id: REQ-API0002-003, description: d, status: pending}\n"
        "  - {id: REQ-API0002-002, description: d, status: pending}\n",
        encoding="utf-8",
    )
    spec = make_spec(
        repo / "spec.xlsx",
        order=False,
        extra_member_rows=[[4, "MemberService", "register", "REQ-1", "", "形式違い", "", "", ""]],
    )
    assert run(repo, str(spec), "--dry-run", "--trace", "traceability/specs") == 0
    out = capsys.readouterr().out
    assert "要件 ID の形式が違います: REQ-1" in out
    assert "REQ-API0001-002 が対応表にありません" in out
    assert "REQ-API0002-003 が対応表にありません" not in out


def test_strict_without_warnings_is_0(repo):
    spec = make_spec(repo / "spec.xlsx")
    assert run(repo, str(spec), "--dry-run", "--strict") == 0


@pytest.mark.parametrize(
    "config, message",
    [
        ("unknown_key: 1\n", "知らないキー"),
        ("header_row: 0\n", "header_row"),
        ("header_row: abc\n", "header_row"),
        ("columns:\n  nothing: 項番\n", "知らない列"),
        ("columns:\n  id: {column: 1A}\n", "列記号"),
        ("columns:\n  id: [No]\n", "引用符で囲む"),
        ("columns:\n  id: {column: A}\n  target_class: {column: A}\n", "同じ列記号"),
        ("requirement_id_pattern: '['\n", "正規表現"),
        ("java:\n  package: Com.Example\n", "java.package"),
        ("java: [1]\n", "java はマッピング"),
        ("sheets: [存在しないシート]\n", "シートがありません"),
        ("sheets: [表紙]\n", "見出しがありません"),
        ("header_row: 2\n", "見出し行が見つかりません"),
        ("- a\n- b\n", "最上位はマッピング"),
        ("key: [unclosed\n", "YAML として読めません"),
    ],
)
def test_invalid_config_exits_2(repo, config, message, capsys):
    spec = make_spec(repo / "spec.xlsx")
    cfg = repo / "bad.yaml"
    cfg.write_text(config, encoding="utf-8")
    assert run(repo, str(spec), "-c", str(cfg), "--dry-run") == 2
    assert message in capsys.readouterr().err


def test_invalid_input_exits_2(repo, capsys):
    assert run(repo, str(repo / "none.xlsx")) == 2
    (repo / "old.xls").write_bytes(b"\xd0\xcf\x11\xe0")
    assert run(repo, str(repo / "old.xls")) == 2
    (repo / "broken.xlsx").write_bytes(b"not a zip")
    assert run(repo, str(repo / "broken.xlsx")) == 2
    assert run(repo) == 2  # 仕様書の指定なし
    assert run(repo, str(make_spec(repo / "s.xlsx")), "-c", str(repo / "none.yaml")) == 2
    err = capsys.readouterr().err
    assert "テスト仕様書がありません" in err and ".xls" in err and "Excel として読めません" in err


def test_unexpected_error_exits_3(repo, monkeypatch, capsys):
    import test_scaffold.cli as cli

    def boom(*_a, **_k):
        raise RuntimeError("boom")

    monkeypatch.setattr(cli, "read_workbook", boom)
    assert run(repo, str(make_spec(repo / "spec.xlsx"))) == 3
    assert "想定外" in capsys.readouterr().err


def test_module_entrypoint_and_print_config():
    r = subprocess.run(
        [sys.executable, "-m", "test_scaffold", "--print-config"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert r.returncode == 0, r.stderr
    loaded = yaml.safe_load(r.stdout)
    assert loaded["header_row"] == "auto" and "項番" in loaded["columns"]["id"]
